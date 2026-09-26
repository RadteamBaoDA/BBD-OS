from __future__ import annotations

import csv
import io
import json
import multiprocessing
import zipfile
from asyncio import to_thread, wait_for
from dataclasses import dataclass
from pathlib import Path
from queue import Empty

from docx import Document as DocxDocument
from pypdf import PdfReader

DOCX_EXPANDED_LIMIT = 100 * 1024 * 1024
PDF_PAGE_LIMIT = 500


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    metadata: dict[str, object]
    warnings: list[str]


def _read_utf8(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _parse_csv(path: Path) -> str:
    rows = list(csv.reader(io.StringIO(_read_utf8(path), newline="")))
    if not rows:
        return ""
    headers = rows[0]
    return "\n\n".join(
        "\n".join(f"{headers[index] if index < len(headers) else f'column_{index + 1}'}: {value}" for index, value in enumerate(row))
        for row in rows[1:]
    ) or ", ".join(headers)


def _parse_pdf(path: Path, page_limit: int) -> ParsedDocument:
    reader = PdfReader(path, strict=True)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported")
    if len(reader.pages) > page_limit:
        raise ValueError("PDF exceeds the configured page limit")
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
    if not text:
        return ParsedDocument("", {"format": "pdf", "pages": len(reader.pages)}, ["No extractable text; OCR is required"])
    return ParsedDocument(text, {"format": "pdf", "pages": len(reader.pages)}, [])


def _parse_docx(path: Path, expanded_limit: int) -> str:
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > expanded_limit:
            raise ValueError("DOCX expanded size exceeds the configured limit")
        names = set(archive.namelist())
        if "word/document.xml" not in names or "[Content_Types].xml" not in names:
            raise ValueError("Invalid DOCX file")
    document = DocxDocument(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
    return "\n".join(parts)


def parse_file(
    path: Path,
    mime: str,
    docx_expanded_limit: int = DOCX_EXPANDED_LIMIT,
    pdf_page_limit: int = PDF_PAGE_LIMIT,
) -> ParsedDocument:
    if mime == "application/pdf":
        return _parse_pdf(path, pdf_page_limit)
    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return ParsedDocument(_parse_docx(path, docx_expanded_limit), {"format": "docx"}, [])
    if mime == "application/json":
        value = json.loads(_read_utf8(path))
        return ParsedDocument(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True), {"format": "json"}, [])
    if mime == "text/csv":
        return ParsedDocument(_parse_csv(path), {"format": "csv"}, [])
    if mime in {"text/plain", "text/markdown"}:
        return ParsedDocument(_read_utf8(path), {"format": "markdown" if mime == "text/markdown" else "txt"}, [])
    raise ValueError("Unsupported file type")


def _parse_process(
    output: multiprocessing.Queue[tuple[str, object]],
    path: Path,
    mime: str,
    docx_expanded_limit: int,
    pdf_page_limit: int,
) -> None:
    try:
        output.put(("ok", parse_file(path, mime, docx_expanded_limit, pdf_page_limit)))
    except Exception as exc:
        output.put((type(exc).__name__, str(exc)))


async def parse_file_bounded(
    path: Path,
    mime: str,
    timeout_seconds: int,
    docx_expanded_limit: int,
    pdf_page_limit: int,
) -> ParsedDocument:
    context = multiprocessing.get_context("spawn")
    output = context.Queue(maxsize=1)
    process = context.Process(
        target=_parse_process,
        args=(output, path, mime, docx_expanded_limit, pdf_page_limit),
    )
    process.start()
    try:
        try:
            status, result = await wait_for(to_thread(output.get, True, timeout_seconds), timeout_seconds + 1)
        except (TimeoutError, Empty) as exc:
            raise TimeoutError("File parser exceeded its time limit") from exc
        if status != "ok":
            raise ValueError(str(result))
        if not isinstance(result, ParsedDocument):
            raise RuntimeError("File parser returned an invalid result")
        return result
    finally:
        if process.is_alive():
            process.terminate()
        await to_thread(process.join)
        output.close()
